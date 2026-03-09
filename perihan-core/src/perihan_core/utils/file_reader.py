import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from PyPDF2 import PdfReader


def read_file_content(file_path: Path) -> str | None:
    """Read content from a supported document file."""
    suffix = file_path.suffix.lower()
    if suffix in (".txt", ".md"):
        return file_path.read_text(encoding="utf-8", errors="ignore")
    elif suffix == ".pdf":
        reader = PdfReader(str(file_path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        doc = Document(str(file_path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table content
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    elif suffix == ".xml":
        tree = ET.parse(str(file_path))
        parts = []
        for el in tree.iter():
            if el.attrib:
                parts.append(" ".join(str(v) for v in el.attrib.values()))
            if el.text and el.text.strip():
                parts.append(el.text.strip())
        return " ".join(parts)
    return None
